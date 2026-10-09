"""
Unit tests for token stats calculation
"""
import pytest
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).parent.parent / 'python'))

import converter as conv


class TestTokenStats:
    """Test token statistics calculation."""
    
    def test_token_stats_basic(self, tmp_path):
        """Test basic token stats calculation."""
        # Create a simple markdown file
        md_file = tmp_path / "test.md"
        md_file.write_text("# Test\n\nThis is a test document.")
        
        src_text = "This is the source text content."
        stats = conv.token_stats(src_text, md_file)
        
        assert 'tiktoken_available' in stats
        assert 'src_tokens' in stats
        assert 'out_tokens' in stats
        assert 'savings_pct' in stats
        assert 'method' in stats
    
    def test_token_stats_with_source_file(self, tmp_path):
        """Test token stats with source file."""
        # Create source and output files
        src_file = tmp_path / "source.txt"
        src_file.write_text("Source content here")
        
        md_file = tmp_path / "output.md"
        md_file.write_text("# Markdown output")
        
        src_text = src_file.read_text()
        stats = conv.token_stats(src_text, md_file, src=src_file)
        
        assert stats['src_tokens'] > 0
        assert stats['out_tokens'] > 0
    
    def test_token_stats_empty_content(self, tmp_path):
        """Test token stats with empty content."""
        md_file = tmp_path / "empty.md"
        md_file.write_text("")
        
        stats = conv.token_stats("", md_file)
        
        assert stats['src_tokens'] == 0
        assert stats['out_tokens'] == 0
    
    def test_token_stats_large_document(self, tmp_path):
        """Test token stats with larger document."""
        md_file = tmp_path / "large.md"
        large_content = "# Title\n\n" + "This is a paragraph. " * 100
        md_file.write_text(large_content)
        
        src_text = "Source text " * 100
        stats = conv.token_stats(src_text, md_file)
        
        assert stats['src_tokens'] > 0
        assert stats['out_tokens'] > 0


class TestTokenCounting:
    """Test token counting with and without tiktoken."""

    def test_count_tokens(self):
        text = "This is a test document."
        tokens = conv._count_tokens(text)
        expected = len(conv._enc.encode(text)) if conv.TIKTOKEN_AVAILABLE else len(text) // 4
        assert tokens == expected

    def test_count_tokens_empty(self):
        assert conv._count_tokens("") == 0

    def test_count_tokens_large(self):
        text = "This is a test. " * 1000
        assert conv._count_tokens(text) > 0


class TestPdfPageExtraction:
    def test_mixed_pdf_uses_page_level_ocr_and_keeps_page_image(self, tmp_path, monkeypatch):
        pytest.importorskip("fitz")
        fitz = conv._require("PyMuPDF", "fitz")
        document = fitz.open()
        text_page = document.new_page()
        text_page.insert_text((72, 72), "This searchable page has enough text to remain on the text extraction path.")
        scanned_page = document.new_page()
        scanned_page.insert_text((72, 72), "tiny")
        pdf_path = tmp_path / "mixed.pdf"
        document.save(str(pdf_path))
        document.close()
        monkeypatch.setattr(conv, "_ocr_image_bytes", lambda image_data: "Recovered scan text")

        output_path, source_text = conv.convert(pdf_path, output_dir=tmp_path / "converted", return_text=True)
        markdown = output_path.read_text(encoding="utf-8")
        image_path = tmp_path / "converted" / "mixed.pdf_images" / "image_001.png"

        assert source_text is not None
        assert "searchable page" in source_text
        assert "Recovered scan text" in source_text
        assert "![Scanned PDF page 2](mixed.pdf_images/image_001.png)" in markdown
        assert image_path.exists()


class TestRagContentPreservation:
    def test_docx_image_is_linked_and_ocr_is_in_markdown(self, tmp_path, monkeypatch):
        docx = pytest.importorskip("docx")
        image_module = pytest.importorskip("PIL.Image")
        from io import BytesIO

        document = docx.Document()
        image_bytes = BytesIO()
        image_module.new("RGB", (2, 2), "white").save(image_bytes, format="PNG")
        document.add_picture(BytesIO(image_bytes.getvalue()))
        source_path = tmp_path / "picture.docx"
        document.save(source_path)
        monkeypatch.setattr(conv, "_ocr_image_bytes", lambda image_data: "Recognized image text")

        output_path = conv.convert(source_path, output_dir=tmp_path / "converted")
        markdown = output_path.read_text(encoding="utf-8")
        asset_path = tmp_path / "converted" / "picture.docx_images" / "image_001.png"

        assert "![Embedded DOCX image" in markdown
        assert "Recognized image text" in markdown
        assert asset_path.exists()

    def test_duplicate_html_images_are_saved_and_ocrd_once(self, tmp_path, monkeypatch):
        import base64

        encoded = base64.b64encode(b"repeat image data").decode("ascii")
        source_path = tmp_path / "repeated.html"
        source_path.write_text(
            f'<img src="data:image/gif;base64,{encoded}" alt="Figure">'
            f'<img src="data:image/gif;base64,{encoded}" alt="Figure">',
            encoding="utf-8",
        )
        monkeypatch.setattr(conv, "_OCR_TEXT_CACHE", {})
        ocr_calls = []
        monkeypatch.setattr(
            conv,
            "_recognize_image_bytes",
            lambda image_data: ocr_calls.append(image_data) or "Repeated figure text",
        )

        output_path = conv.convert(source_path, output_dir=tmp_path / "converted")
        markdown = output_path.read_text(encoding="utf-8")
        assets = list((tmp_path / "converted" / "repeated.html_images").glob("*"))

        assert len(assets) == 1
        assert markdown.count("repeated.html_images/image_001.gif") == 2
        assert markdown.count("Repeated figure text") == 2
        assert len(ocr_calls) == 1

    def test_embedded_image_ocr_runs_concurrently(self, tmp_path, monkeypatch):
        import base64
        import threading

        barrier = threading.Barrier(4)

        def recognize(image_data):
            barrier.wait(timeout=5)
            return f"OCR result {image_data[-1]}"

        source_path = tmp_path / "images.html"
        image_tags = [
            f'<img src="data:image/gif;base64,{base64.b64encode(bytes([i])).decode()}" alt="Image">'
            for i in range(1, 5)
        ]
        source_path.write_text("".join(image_tags), encoding="utf-8")
        monkeypatch.setattr(conv, "_OCR_TEXT_CACHE", {})
        monkeypatch.setattr(conv, "_recognize_image_bytes", recognize)

        output_path = conv.convert(source_path, output_dir=tmp_path / "converted")
        markdown = output_path.read_text(encoding="utf-8")

        for index in range(1, 5):
            assert f"OCR result {index}" in markdown

    def test_pptx_chart_data_and_picture_are_preserved(self, tmp_path, monkeypatch):
        pptx = pytest.importorskip("pptx")
        image_module = pytest.importorskip("PIL.Image")
        from io import BytesIO
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2"]
        chart_data.add_series("Revenue", [10, 20])
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1), Inches(1), Inches(5), Inches(3), chart_data,
        )
        image_bytes = BytesIO()
        image_module.new("RGB", (2, 2), "white").save(image_bytes, format="PNG")
        slide.shapes.add_picture(BytesIO(image_bytes.getvalue()), Inches(1), Inches(4))
        source_path = tmp_path / "charts.pptx"
        presentation.save(source_path)
        monkeypatch.setattr(conv, "_ocr_image_bytes", lambda image_data: "Chart image labels")

        output_path = conv.convert(source_path, output_dir=tmp_path / "converted")
        markdown = output_path.read_text(encoding="utf-8")
        asset_path = tmp_path / "converted" / "charts.pptx_images" / "image_001.png"

        assert "| Category | Revenue |" in markdown
        assert "| Q1 | 10 |" in markdown
        assert "| Q2 | 20 |" in markdown
        assert "![Image on slide 1]" in markdown
        assert "Chart image labels" in markdown
        assert asset_path.exists()

    def test_xlsx_preserves_formula_and_cached_value(self, tmp_path, monkeypatch):
        openpyxl = pytest.importorskip("openpyxl")
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Metrics"
        worksheet.append(["Name", "Value", "Double"])
        worksheet.append(["A", 5, "=B2*2"])
        source_path = tmp_path / "formulas.xlsx"
        workbook.save(source_path)
        uncached_output = conv._convert_excel(source_path)
        assert "=B2*2 [cached value unavailable]" in uncached_output
        load_workbook = openpyxl.load_workbook

        def load_with_cached_formula_value(path, data_only=False):
            loaded = load_workbook(path, data_only=data_only)
            if data_only:
                loaded["Metrics"]["C2"] = 10
            return loaded

        monkeypatch.setattr(openpyxl, "load_workbook", load_with_cached_formula_value)
        output = conv._convert_excel(source_path)

        assert "10 [formula: =B2*2]" in output
        assert "| A | 5 | 10 [formula: =B2*2] |" in output


class TestFriendlyErrorMessages:
    def test_missing_python_package_suggests_requirements(self):
        exc = ImportError("Required package 'PyMuPDF' is not installed.")
        message = conv.friendly_error_message(exc)
        assert "PyMuPDF" in message
        assert str(Path(conv.__file__).with_name("requirements.txt")) in message

    def test_missing_tesseract_suggests_install_and_path(self):
        exc = type("TesseractNotFoundError", (Exception,), {})()
        message = conv.friendly_error_message(exc)
        assert "Tesseract" in message
        assert "PATH" in message

    def test_permission_error_suggests_writable_folder(self):
        message = conv.friendly_error_message(PermissionError("access denied"))
        assert "write permission" in message


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
